#!/usr/bin/env python3
"""在课程模板中生成作业要求对应的报告，直接嵌入浏览器原始截图。

正文取自 docs/assignments/hwX/report.md。保留模板的页面与页脚，替换表头和正文槽位；
图片按原始比例插入，不生成、重绘或拼接截图。
"""

from __future__ import annotations

import os
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[2]
TEMPLATE = ROOT / "scripts/homework/templates/作业模版.docx"
STUDENT_ID = "524031910113"
STUDENT_NAME = "俞冠廷"


def font_settings(font, properties, size: float, latin: str = "Times New Roman") -> None:
    """使用可显示中文的字体，避免模板的宋体在本机渲染时缺字。"""
    font.name = latin
    font.size = Pt(size)
    font.color.rgb = RGBColor(0, 0, 0)
    fonts = properties.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        properties.insert(0, fonts)
    fonts.set(qn("w:eastAsia"), "Source Han Sans CN")


def plain(text: str) -> str:
    """将 Markdown 行内代码标记转为正文，不引入额外装饰格式。"""
    return text.replace("`", "")


def format_paragraph(paragraph, *, code: bool = False) -> None:
    """正文采用自然段间距；代码保留逐行缩进并使用等宽字体。"""
    form = paragraph.paragraph_format
    form.left_indent = Cm(0)
    form.right_indent = Cm(0)
    form.first_line_indent = Cm(0)
    form.space_before = Pt(0)
    form.space_after = Pt(6)
    form.line_spacing = Pt(14.5 if not code else 11)
    form.widow_control = True
    properties = paragraph._element.get_or_add_pPr()
    grid = properties.find(qn("w:snapToGrid"))
    if grid is None:
        grid = OxmlElement("w:snapToGrid")
        properties.append(grid)
    grid.set(qn("w:val"), "0")
    if code:
        form.keep_together = True
    for run in paragraph.runs:
        font_settings(run.font, run._element.get_or_add_rPr(), 8.5 if code else 11,
                      "Menlo" if code else "Times New Roman")


def add_table(document, lines: list[str]) -> None:
    """将接口映射表写入文档，明确设置边框与表头。"""
    rows = [
        [plain(cell.strip()) for cell in line.strip().strip("|").split("|")]
        for line in lines
        if not re.match(r"^\|[\s:|-]+\|$", line)
    ]
    table = document.add_table(rows=0, cols=len(rows[0]))
    table.autofit = False
    widths = [Cm(1.8), Cm(5.2), Cm(4.2), Cm(3.25)]
    if rows[0][0] == "函数":
        widths = [Cm(4.2), Cm(2.0), Cm(4.6), Cm(3.65)]
    for column, width in zip(table.columns, widths):
        column.width = width
    for index, values in enumerate(rows):
        cells = table.add_row().cells
        for column, value in enumerate(values):
            cell = cells[column]
            cell.width = widths[column]
            cell.text = value
            props = cell._tc.get_or_add_tcPr()
            borders = OxmlElement("w:tcBorders")
            for edge in ("top", "left", "bottom", "right"):
                item = OxmlElement(f"w:{edge}")
                item.set(qn("w:val"), "single")
                item.set(qn("w:sz"), "4")
                item.set(qn("w:color"), "D9D9D9")
                borders.append(item)
            props.append(borders)
            if index == 0:
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), "EEEEEE")
                props.append(shading)
            for paragraph in cell.paragraphs:
                format_paragraph(paragraph)
                paragraph.paragraph_format.space_after = Pt(4)
                paragraph.paragraph_format.space_before = Pt(4)
                paragraph.paragraph_format.line_spacing = Pt(12)
                for run in paragraph.runs:
                    font_settings(run.font, run._element.get_or_add_rPr(), 9)
    header_props = table.rows[0]._tr.get_or_add_trPr()
    header_props.append(OxmlElement("w:tblHeader"))


def build_report(source: Path, output: Path, number: int) -> None:
    """填入模板表头，解析五节 Markdown，插入实际截图后保存。"""
    document = Document(TEMPLATE)
    for paragraph in list(document.paragraphs)[4:]:
        paragraph._element.getparent().remove(paragraph._element)
    from datetime import datetime
    from zoneinfo import ZoneInfo
    document.paragraphs[0].text = datetime.now(ZoneInfo("Asia/Shanghai")).strftime("%Y-%m-%d")
    document.paragraphs[1].text = f"应用系统体系架构 — 作业{number}"
    document.paragraphs[2].text = f"学号：{STUDENT_ID}    姓名：{STUDENT_NAME}    得分：________"
    document.paragraphs[3].text = "本次作业回答如下："
    for style in document.styles:
        if style.type == 1:
            props = style.element.get_or_add_pPr()
            for numbering in list(props.findall(qn("w:numPr"))):
                props.remove(numbering)
            font_settings(style.font, style.element.get_or_add_rPr(), 11)
    for index, paragraph in enumerate(document.paragraphs):
        format_paragraph(paragraph)
        paragraph.paragraph_format.keep_with_next = True
        if index == 1:
            paragraph.paragraph_format.line_spacing = Pt(22)
            for run in paragraph.runs:
                font_settings(run.font, run._element.get_or_add_rPr(), 15)
    lines = source.read_text(encoding="utf-8").splitlines()
    cursor = 0
    while cursor < len(lines):
        line = lines[cursor]
        if not line.strip():
            cursor += 1
            continue
        if line.startswith("## "):
            title = line[3:]
            paragraph = document.add_heading(title, level=2)
            format_paragraph(paragraph)
            paragraph.paragraph_format.keep_with_next = True
            paragraph.paragraph_format.space_before = Pt(12)
            paragraph.paragraph_format.line_spacing = Pt(18)
            for run in paragraph.runs:
                font_settings(run.font, run._element.get_or_add_rPr(), 13)
        elif line.startswith("```"):
            code = []
            cursor += 1
            while cursor < len(lines) and not lines[cursor].startswith("```"):
                code.append(lines[cursor])
                cursor += 1
            paragraph = document.add_paragraph("\n".join(code))
            format_paragraph(paragraph, code=True)
        elif line.startswith("|"):
            table_lines = []
            while cursor < len(lines) and lines[cursor].startswith("|"):
                table_lines.append(lines[cursor])
                cursor += 1
            add_table(document, table_lines)
            continue
        elif line.startswith("!["):
            match = re.fullmatch(r"!\[(.*?)\]\((.*?)\)", line)
            if match is None:
                raise ValueError(f"图片引用格式错误：{line}")
            paragraph = document.add_paragraph()
            document.paragraphs[-2].paragraph_format.keep_with_next = True
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.keep_with_next = True
            paragraph.add_run().add_picture(str(source.parent / match[2]), width=Cm(14.2))
            caption = document.add_paragraph(match[1])
            format_paragraph(caption)
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in caption.runs:
                run.font.size = Pt(9)
        else:
            paragraph = document.add_paragraph(plain(line))
            format_paragraph(paragraph)
        cursor += 1
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    print(f"已生成 {output.name}")

